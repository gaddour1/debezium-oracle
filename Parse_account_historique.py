"""
Pipeline PySpark — Parser XML T24 (champs simples + multiples)
JOB HISTORIQUE / BATCH
====================================================================
Source  : Kafka topic biat.C__DBZUSER.ACCOUNT
Mode    : Batch (lit tout l'historique avec earliest)
Output  : Iceberg / MinIO (tables persistées)
"""

import sys
sys.path.insert(0, '/opt/spark-apps/libs')

from pyspark.sql import SparkSession
from pyspark.sql.functions import col, get_json_object, udf, explode
from pyspark.sql.types import MapType, StringType, ArrayType, StructType, StructField
import xml.etree.ElementTree as ET
import base64
import openpyxl


# =============================================================================
# ETAPE 1 - Charger le dictionnaire depuis Excel (dynamique)
# =============================================================================
def load_dictionary_dynamic(xlsx_path, application="ACCOUNT"):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(cell.value).strip() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx_app   = headers.index('APPLICATION')
    idx_ordre = headers.index('ORDRE_CHAMPS')
    idx_champ = headers.index('CHAMP')
    idx_multi = headers.index('CHAMP_MUTIPLE')
    column_map_simple   = {}
    column_map_multiple = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        app   = str(row[idx_app]).strip().upper()  if row[idx_app]   else ""
        ordre = str(row[idx_ordre]).strip()         if row[idx_ordre] else ""
        champ = str(row[idx_champ]).strip()         if row[idx_champ] else ""
        multi = str(row[idx_multi]).strip().upper() if row[idx_multi] else ""
        if app == application.upper() and ordre.isdigit():
            champ_norm = champ.replace('.', '_')
            if multi == 'S':
                column_map_simple[f"c{ordre}"] = champ_norm
            elif multi == 'M':
                column_map_multiple[ordre] = champ_norm
    wb.close()
    print(f"Dictionnaire simple   : {len(column_map_simple)} colonnes")
    print(f"Dictionnaire multiple : {len(column_map_multiple)} colonnes")
    return column_map_simple, column_map_multiple


# =============================================================================
# ETAPE 2 - Décoder RECID
# =============================================================================
def decode_recid(b64_value):
    if b64_value is None:
        return None
    try:
        # Essayer de décoder en base64
        raw = base64.b64decode(b64_value, validate=True)
        # Vérifier que le résultat est un nombre (RECID Oracle est numérique)
        decoded = str(int.from_bytes(raw, byteorder='big', signed=False))
        # Si le résultat est 0 ou absurde, c'était pas du base64
        if decoded == '0':
            return b64_value
        return decoded
    except Exception:
        # Pas du base64 → retourner la valeur telle quelle
        return b64_value


# =============================================================================
# ETAPE 3A - UDF parse XML simple
# =============================================================================
def make_parse_simple_udf(column_map_simple):
    col_map = dict(column_map_simple)
    def parse_simple(xml_str):
        if not xml_str:
            return {}
        try:
            root = ET.fromstring(xml_str.strip())
            result = {}
            for child in root:
                if child.get('m') is None:
                    col_name = col_map.get(child.tag)
                    if col_name:
                        result[col_name] = (child.text or '').strip()
            return result
        except ET.ParseError:
            return {}
    return udf(parse_simple, MapType(StringType(), StringType()))


# =============================================================================
# ETAPE 3B - UDF parse XML multiple
# =============================================================================
MULTI_SCHEMA = ArrayType(StructType([
    StructField("CHAMP",  StringType(), True),
    StructField("BLOC",   StringType(), True),
    StructField("SUB",    StringType(), True),
    StructField("VALEUR", StringType(), True),
]))

def make_parse_multiple_udf(column_map_multiple):
    col_map = dict(column_map_multiple)
    def parse_multiple(xml_str):
        if not xml_str:
            return []
        try:
            root = ET.fromstring(xml_str.strip())
            result = []
            for child in root:
                tag = child.tag
                m   = child.get('m')
                s   = child.get('s')
                val = (child.text or '').strip()
                ordre = ''.join(filter(str.isdigit, tag))
                champ = col_map.get(ordre)
                if champ:
                    bloc = m if m else '1'
                    sub  = s if s else '1'
                    result.append((champ, bloc, sub, val))
            return result
        except ET.ParseError:
            return []
    return udf(parse_multiple, MULTI_SCHEMA)


# =============================================================================
# JOB HISTORIQUE - BATCH
# =============================================================================
def run_historique():

    # ── SparkSession avec Iceberg + MinIO ────────────────────────────────────
    spark = SparkSession.builder \
        .appName("T24-Account-Historique-Batch") \
        .master("local[*]") \
        .config("spark.sql.extensions",
            "org.apache.iceberg.spark.extensions.IcebergSparkSessionExtensions") \
        .config("spark.sql.catalog.biat",
            "org.apache.iceberg.spark.SparkCatalog") \
        .config("spark.sql.catalog.biat.type", "rest") \
        .config("spark.sql.catalog.biat.uri", "http://iceberg-rest:8181") \
        .config("spark.sql.catalog.biat.warehouse", "s3://warehouse/") \
        .config("spark.sql.catalog.biat.io-impl",
            "org.apache.iceberg.aws.s3.S3FileIO") \
        .config("spark.sql.catalog.biat.s3.endpoint", "http://minio:9000") \
        .config("spark.sql.catalog.biat.s3.path-style-access", "true") \
        .config("spark.sql.catalog.biat.s3.access-key-id", "minioadmin") \
        .config("spark.sql.catalog.biat.s3.secret-access-key", "minioadmin") \
        .config("spark.sql.catalog.biat.s3.region", "us-east-1") \
        .config("spark.hadoop.fs.s3a.endpoint", "http://minio:9000") \
        .config("spark.hadoop.fs.s3a.access.key", "minioadmin") \
        .config("spark.hadoop.fs.s3a.secret.key", "minioadmin") \
        .config("spark.hadoop.fs.s3a.path.style.access", "true") \
        .config("spark.hadoop.fs.s3a.impl",
            "org.apache.hadoop.fs.s3a.S3AFileSystem") \
        .getOrCreate()

    spark.sparkContext.setLogLevel("WARN")

    # ── ETAPE 1 : Créer le namespace ─────────────────────────────────────────
    spark.sql("CREATE NAMESPACE IF NOT EXISTS biat.t24")
    print("Namespace biat.t24 prêt")

    # ── ETAPE 2 : Charger le dictionnaire ────────────────────────────────────
    # On charge le dictionnaire AVANT de créer les tables
    # car on a besoin des noms de colonnes pour créer account_simple
    column_map_simple, column_map_multiple = load_dictionary_dynamic(
        "/opt/spark-apps/DICTIONNAIRE_T24.xlsx", "ACCOUNT"
    )
    if not column_map_simple:
        print("Dictionnaire simple vide — arrêt")
        return

    col_names = sorted(set(column_map_simple.values()))

    # ── ETAPE 3 : Créer les tables Iceberg avec le bon schéma ────────────────
    # On supprime si elles existent déjà (pour éviter le problème de schéma)
    spark.sql("DROP TABLE IF EXISTS biat.t24.account_simple")
    spark.sql("DROP TABLE IF EXISTS biat.t24.account_multiple")

    # account_simple : 3 colonnes fixes + toutes les colonnes du dictionnaire
    # Si un compte n'a pas de valeur pour une colonne → NULL automatique
    colonnes_simple = ",\n        ".join([f"`{nom}` STRING" for nom in col_names])
    spark.sql(f"""
        CREATE TABLE biat.t24.account_simple (
            RECID    STRING,
            op       STRING,
            kafka_ts TIMESTAMP,
            {colonnes_simple}
        )
        USING iceberg
        PARTITIONED BY (days(kafka_ts))
    """)

    # account_multiple : schéma fixe — les valeurs sont en lignes pas en colonnes
    spark.sql("""
        CREATE TABLE biat.t24.account_multiple (
            RECID    STRING,
            op       STRING,
            kafka_ts TIMESTAMP,
            CHAMP    STRING,
            BLOC     STRING,
            SUB      STRING,
            VALEUR   STRING
        )
        USING iceberg
        PARTITIONED BY (days(kafka_ts))
    """)

    print(f"Table account_simple créée avec {len(col_names) + 3} colonnes")
    print("Table account_multiple créée avec 7 colonnes")

    # ── ETAPE 4 : UDFs ───────────────────────────────────────────────────────
    parse_simple_udf   = make_parse_simple_udf(column_map_simple)
    parse_multiple_udf = make_parse_multiple_udf(column_map_multiple)
    decode_recid_udf   = udf(decode_recid, StringType())

    # ── ETAPE 5 : Lecture Kafka batch ────────────────────────────────────────
    print("Lecture de l'historique complet depuis Kafka...")
    df_kafka = spark.read \
        .format("kafka") \
        .option("kafka.bootstrap.servers", "kafka:9092") \
        .option("subscribe", "biat.C__DBZUSER.ACCOUNT") \
        .option("startingOffsets", "earliest") \
        .load()

    df_raw = df_kafka.select(
        get_json_object(col("value").cast("string"), "$.payload.after.RECID").alias("recid_b64"),
        get_json_object(col("value").cast("string"), "$.payload.after.XMLRECORD").alias("xmlrecord"),
        get_json_object(col("value").cast("string"), "$.payload.op").alias("op"),
        col("timestamp").alias("kafka_ts")
    ).filter(col("xmlrecord").isNotNull())

    df_base = df_raw.withColumn("RECID", decode_recid_udf(col("recid_b64")))

    # ── ETAPE 6 : Parsing XML ────────────────────────────────────────────────
    df_simple_final = df_base \
        .withColumn("fields", parse_simple_udf(col("xmlrecord"))) \
        .select(
            col("RECID"), col("op"), col("kafka_ts"),
            *[col("fields")[nom].alias(nom) for nom in col_names]
        )

    df_multiple_final = df_base \
        .withColumn("multi_rows", parse_multiple_udf(col("xmlrecord"))) \
        .withColumn("multi", explode(col("multi_rows"))) \
        .select(
            col("RECID"), col("op"), col("kafka_ts"),
            col("multi.CHAMP").alias("CHAMP"),
            col("multi.BLOC").alias("BLOC"),
            col("multi.SUB").alias("SUB"),
            col("multi.VALEUR").alias("VALEUR")
        )

    # ── ETAPE 7 : Écriture dans Iceberg ──────────────────────────────────────
    print("Écriture dans Iceberg...")

    df_simple_final.writeTo("biat.t24.account_simple").append()
    print(f"Table simple   écrite : {df_simple_final.count()} lignes")

    df_multiple_final.writeTo("biat.t24.account_multiple").append()
    print(f"Table multiple écrite : {df_multiple_final.count()} lignes")

    # ── ETAPE 8 : Vérification ───────────────────────────────────────────────
    print("\nAperçu table simple dans Iceberg :")
    spark.sql("SELECT RECID, op, kafka_ts, CURRENCY, CUSTOMER FROM biat.t24.account_simple LIMIT 5").show(truncate=False)

    print("\nAperçu table multiple dans Iceberg :")
    spark.sql("SELECT * FROM biat.t24.account_multiple LIMIT 10").show(truncate=False)

    print("\nJob Batch Historique terminé avec succès !")
    spark.stop()


if __name__ == "__main__":
    run_historique()