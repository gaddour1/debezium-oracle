"""
Pipeline PySpark — Parser XML T24 (champs simples + multiples)
JOB STREAMING (temps réel)
====================================================================
Source  : Kafka topic biat.C__DBZUSER.ACCOUNT
Mode    : Streaming
Output  : Iceberg / MinIO
"""

import sys
sys.path.insert(0, '/opt/spark-apps/libs')

from pyspark.sql import SparkSession
from pyspark.sql.functions import col, get_json_object, udf, explode
from pyspark.sql.types import MapType, StringType, ArrayType, StructType, StructField
import xml.etree.ElementTree as ET
import base64
import openpyxl


# =============================================================================
# 1. Charger le dictionnaire depuis Excel
# =============================================================================
def load_dictionary_dynamic(xlsx_path, application="ACCOUNT"):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(cell.value).strip() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    
    idx_app   = headers.index('APPLICATION')
    idx_ordre = headers.index('ORDRE_CHAMPS')
    idx_champ = headers.index('CHAMP')
    idx_multi = headers.index('CHAMP_MUTIPLE')

    column_map_simple   = {}
    column_map_multiple = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        app   = str(row[idx_app]).strip().upper()   if row[idx_app]   else ""
        ordre = str(row[idx_ordre]).strip()          if row[idx_ordre] else ""
        champ = str(row[idx_champ]).strip()          if row[idx_champ] else ""
        multi = str(row[idx_multi]).strip().upper()  if row[idx_multi] else ""

        if app == application.upper() and ordre.isdigit():
            champ_norm = champ.replace('.', '_')
            if multi == 'S':
                column_map_simple[f"c{ordre}"] = champ_norm
            elif multi == 'M':
                column_map_multiple[ordre] = champ_norm

    wb.close()
    print(f"Dictionnaire simple   : {len(column_map_simple)} colonnes")
    print(f"Dictionnaire multiple : {len(column_map_multiple)} colonnes")
    return column_map_simple, column_map_multiple


# =============================================================================
# 2. Decode RECID
# =============================================================================
def decode_recid(value):
    if value is None:
        return None
    try:
        raw = base64.b64decode(value, validate=True)
        decoded = str(int.from_bytes(raw, byteorder='big', signed=False))
        if decoded == '0' or len(decoded) < 3:
            return value
        return decoded
    except Exception:
        return value


# =============================================================================
# 3A. Parse XML Simple
# =============================================================================
def make_parse_simple_udf(column_map_simple):
    col_map = dict(column_map_simple)
    def parse_simple(xml_str):
        if not xml_str:
            return {}
        try:
            root = ET.fromstring(xml_str.strip())
            result = {}
            for child in root:
                if child.get('m') is None:
                    col_name = col_map.get(child.tag)
                    if col_name:
                        result[col_name] = (child.text or '').strip()
            return result
        except Exception:
            return {}
    return udf(parse_simple, MapType(StringType(), StringType()))


# =============================================================================
# 3B. Parse XML Multiple
# =============================================================================
MULTI_SCHEMA = ArrayType(StructType([
    StructField("CHAMP",  StringType(), True),
    StructField("BLOC",   StringType(), True),
    StructField("SUB",    StringType(), True),
    StructField("VALEUR", StringType(), True),
]))

def make_parse_multiple_udf(column_map_multiple):
    col_map = dict(column_map_multiple)
    def parse_multiple(xml_str):
        if not xml_str:
            return []
        try:
            root = ET.fromstring(xml_str.strip())
            result = []
            for child in root:
                tag = child.tag
                m   = child.get('m')
                s   = child.get('s')
                val = (child.text or '').strip()
                ordre = ''.join(filter(str.isdigit, tag))
                champ = col_map.get(ordre)
                if champ:
                    bloc = m if m else '1'
                    sub  = s if s else '1'
                    result.append((champ, bloc, sub, val))
            return result
        except Exception:
            return []
    return udf(parse_multiple, MULTI_SCHEMA)


# =============================================================================
# MAIN
# =============================================================================
def run_streaming():
    spark = SparkSession.builder \
        .appName("T24-Account-Streaming") \
        .master("local[*]") \
        .config("spark.sql.extensions", "org.apache.iceberg.spark.extensions.IcebergSparkSessionExtensions") \
        .config("spark.sql.catalog.biat", "org.apache.iceberg.spark.SparkCatalog") \
        .config("spark.sql.catalog.biat.type", "rest") \
        .config("spark.sql.catalog.biat.uri", "http://iceberg-rest:8181") \
        .config("spark.sql.catalog.biat.warehouse", "s3://warehouse/") \
        .config("spark.sql.catalog.biat.io-impl", "org.apache.iceberg.aws.s3.S3FileIO") \
        .config("spark.sql.catalog.biat.s3.endpoint", "http://minio:9000") \
        .config("spark.sql.catalog.biat.s3.path-style-access", "true") \
        .config("spark.sql.catalog.biat.s3.access-key-id", "minioadmin") \
        .config("spark.sql.catalog.biat.s3.secret-access-key", "minioadmin") \
        .config("spark.sql.catalog.biat.s3.region", "us-east-1") \
        .getOrCreate()

    spark.sparkContext.setLogLevel("WARN")

    spark.sql("CREATE NAMESPACE IF NOT EXISTS biat.t24")

    column_map_simple, column_map_multiple = load_dictionary_dynamic(
        "/opt/spark-apps/DICTIONNAIRE_T24.xlsx", "ACCOUNT"
    )

    col_names = sorted(set(column_map_simple.values()))

    # ✅ CREATE TABLE (sans DROP)
    colonnes_simple = ",\n        ".join([f"`{nom}` STRING" for nom in col_names])

    spark.sql(f"""
        CREATE TABLE IF NOT EXISTS biat.t24.account_simple (
            RECID STRING,
            op STRING,
            kafka_ts TIMESTAMP,
            {colonnes_simple}
        )
        USING iceberg
        PARTITIONED BY (days(kafka_ts))
    """)

    spark.sql("""
        CREATE TABLE IF NOT EXISTS biat.t24.account_multiple (
            RECID STRING,
            op STRING,
            kafka_ts TIMESTAMP,
            CHAMP STRING,
            BLOC STRING,
            SUB STRING,
            VALEUR STRING
        )
        USING iceberg
        PARTITIONED BY (days(kafka_ts))
    """)

    parse_simple_udf   = make_parse_simple_udf(column_map_simple)
    parse_multiple_udf = make_parse_multiple_udf(column_map_multiple)
    decode_recid_udf   = udf(decode_recid, StringType())

    df_stream = spark.readStream \
        .format("kafka") \
        .option("kafka.bootstrap.servers", "kafka:9092") \
        .option("subscribe", "biat.C__DBZUSER.ACCOUNT") \
        .option("startingOffsets", "latest") \
        .option("failOnDataLoss", "false") \
        .load()

    df_raw = df_stream.select(
        get_json_object(col("value").cast("string"), "$.payload.after.RECID").alias("recid_b64"),
        get_json_object(col("value").cast("string"), "$.payload.after.XMLRECORD").alias("xmlrecord"),
        get_json_object(col("value").cast("string"), "$.payload.op").alias("op"),
        col("timestamp").alias("kafka_ts")
    ).filter(col("xmlrecord").isNotNull())

    df_base = df_raw.withColumn("RECID", decode_recid_udf(col("recid_b64")))

    df_simple = df_base.withColumn("fields", parse_simple_udf(col("xmlrecord"))).select(
        col("RECID"), col("op"), col("kafka_ts"),
        *[col("fields")[nom].alias(nom) for nom in col_names]
    )

    df_multiple = df_base.withColumn("multi_rows", parse_multiple_udf(col("xmlrecord"))) \
        .withColumn("multi", explode(col("multi_rows"))) \
        .select(
            col("RECID"), col("op"), col("kafka_ts"),
            col("multi.CHAMP"), col("multi.BLOC"),
            col("multi.SUB"), col("multi.VALEUR")
        )

    # ✅ CHECKPOINT FIXÉ
    query_simple = df_simple.writeStream \
        .foreachBatch(lambda df, _: df.writeTo("biat.t24.account_simple").append() if not df.isEmpty() else None) \
        .option("checkpointLocation", "/tmp/checkpoint/account_simple") \
        .start()

    query_multiple = df_multiple.writeStream \
        .foreachBatch(lambda df, _: df.writeTo("biat.t24.account_multiple").append() if not df.isEmpty() else None) \
        .option("checkpointLocation", "/tmp/checkpoint/account_multiple") \
        .start()

    print("✅ Streaming démarré avec succès !")

    query_simple.awaitTermination()
    query_multiple.awaitTermination()


if __name__ == "__main__":
    run_streaming()